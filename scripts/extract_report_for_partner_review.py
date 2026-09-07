#!/usr/bin/env python3
"""
Extracts a completed fm-validator report's real Issue Log and Validation
Matrix into clean JSON, for handoff to Partner Review's Phase B.

Usage:
    python3 scripts/extract_report_for_partner_review.py /path/to/report_VALIDATED.xlsx > report.json
"""
import sys
import json
import openpyxl


def find_header_row(ws, expected_header_text, max_scan=10, max_col_scan=15):
    """Real header rows in fm-validator's own reports aren't always at a
    fixed row number (title/summary rows above them vary), and the header
    text itself isn't always in column 1 either (a narrow, blank
    visual-spacing column often precedes it) - scan both dimensions for
    a genuine match rather than assume a fixed position.
    Returns (row_num, col_num) or (None, None) if not found."""
    for row_num in range(1, max_scan + 1):
        for col_num in range(1, max_col_scan + 1):
            if ws.cell(row=row_num, column=col_num).value == expected_header_text:
                return row_num, col_num
    return None, None


def extract_issue_log(wb):
    ws = wb['Issue Log']
    header_row, id_col = find_header_row(ws, 'ID')
    if header_row is None:
        return []
    headers = [ws.cell(row=header_row, column=c).value for c in range(id_col, ws.max_column + 1)]
    issues = []
    for row_num in range(header_row + 1, ws.max_row + 1):
        row_id = ws.cell(row=row_num, column=id_col).value
        if not row_id:
            continue
        entry = {}
        for offset, header in enumerate(headers):
            if header:
                entry[header] = ws.cell(row=row_num, column=id_col + offset).value
        issues.append(entry)
    return issues


def extract_validation_matrix(wb):
    ws = wb['Validation Matrix']
    header_row, id_col = find_header_row(ws, 'Rule ID')
    if header_row is None:
        return []
    headers = [ws.cell(row=header_row, column=c).value for c in range(id_col, ws.max_column + 1)]
    rules = []
    for row_num in range(header_row + 1, ws.max_row + 1):
        rule_id = ws.cell(row=row_num, column=id_col).value
        if not rule_id:
            continue
        entry = {}
        for offset, header in enumerate(headers):
            if header:
                entry[header] = ws.cell(row=row_num, column=id_col + offset).value
        rules.append(entry)
    return rules


def main():
    if len(sys.argv) < 2:
        print('Usage: python3 scripts/extract_report_for_partner_review.py /path/to/report_VALIDATED.xlsx > report.json', file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    result = {
        'issueLog': extract_issue_log(wb),
        'validationMatrix': extract_validation_matrix(wb),
    }
    print(f"Extracted {len(result['issueLog'])} issue log entries, {len(result['validationMatrix'])} validation matrix rows.", file=sys.stderr)
    # The real extraction goes to stdout, so it can be redirected cleanly
    # without the diagnostic line above getting mixed in.
    # Some columns (e.g. a date-identified field) genuinely contain real
    # datetime objects, not strings - json.dumps can't serialize these
    # natively.
    print(json.dumps(result, default=str))


if __name__ == '__main__':
    main()
